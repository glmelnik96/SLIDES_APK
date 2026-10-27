"""Краска фигур: прозрачность, рамка по сторонам, полосатый фон, перенос слов.

Зачем отдельный модуль. Четыре поломки ниже объединяет одно: ни одну из них не
видит ни fidelity, ни тесты пакета. Fidelity сверяет ЧИСЛА — x, y, w, h рамок, и
на всех четырёх он давал ровные нули: фигура стояла точно на месте, просто
залита была не тем. Тесты пакета читают оболочку — тему, мастер, имена. Поймать
такое можно было только одним способом: открыть файл в PowerPoint и посмотреть
глазами. Так и нашлось — по снимкам, где «почти на каждом слайде что-то едет».

  * ПРОЗРАЧНОСТЬ. `fill.transparency = 0.86` в python-pptx 1.0.2 не существует:
    присваивание молча заводит посторонний атрибут, в файл уходит глухой цвет.
    Зелёный wash под линией line-chart (fill-opacity .14) приезжал СПЛОШНОЙ
    зелёной заливкой во всю высоту графика.
  * РАМКА. У фигуры PPTX обводка одна на весь контур, в CSS — поштучная.
    `border-top:1px` разделителя .svc-row превращался в зелёный ПРЯМОУГОЛЬНИК
    вокруг каждой строки.
  * ПОЛОСАТЫЙ ФОН. «Шурешки» обложки заданы repeating-linear-gradient, а фон
    читался только из backgroundColor — на слайде не появлялось ничего.
  * ПЕРЕНОС. Chrome при overflow-wrap:normal слишком широкое слово выпускает за
    край, PowerPoint рвёт его посреди слова: «СЕТЬ» -> «СЕТ» + «Ь». Рамок,
    у которых ширина совпадает с шириной текста, в эталонной деке 136 из 678.

Единицы проверяются на синтетике (стиль — обычный dict, как в палитре), а потом
один раз на настоящей деке: синтетика говорит, что функция считает верно, дека —
что её кто-то зовёт.
"""
import os
import zipfile
from pathlib import Path

os.environ["SLIDES_APP_SKIP_SHIM"] = "1"

import pytest
from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Pt

from htmlslides.topptx import convert
from htmlslides.topptx.pptxshell import set_alpha

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"

# Фон .cover-hatch в том виде, в каком его отдаёт getComputedStyle — строка
# взята из палитры geometry.json (стиль 8, тёмная тема), а не написана от руки:
# разбирать мы обязаны именно то, что приходит из браузера.
HATCH_BG = ("repeating-linear-gradient(90deg, rgb(38, 208, 124) 0px, "
            "rgb(38, 208, 124) 2px, rgba(0, 0, 0, 0) 2px, "
            "rgba(0, 0, 0, 0) 20px)")
DOTS_BG = "radial-gradient(circle, rgb(38, 208, 124) 1.5px, rgba(0, 0, 0, 0) 1.6px)"


@pytest.fixture(scope="module")
def canvas():
    """Пустой слайд и строитель. Браузер не поднимается: таблица не нужна."""
    b = convert._Builder("dark")
    layout = b.shell.layout_for("тест", "Тест", None)
    return b, b.prs.slides.add_slide(layout)


def _xml(shape) -> str:
    return etree.tostring(shape.element, encoding="unicode")


def _alpha_of(shape) -> int | None:
    """Значение <a:alpha> в заливке фигуры. None, если альфы нет."""
    got = shape.element.findall(f".//{A}solidFill/{A}srgbClr/{A}alpha")
    return int(got[0].get("val")) if got else None


# ------------------------------------------------------------- прозрачность
def test_python_pptx_still_has_no_transparency_property():
    """Проверка предпосылки: свойства нет, значит set_alpha не лишний.

    Тест-страховка от «упрощения». Выглядит `fill.transparency = 0.14`
    настолько естественно, что к нему хочется вернуться, а стоило это худшего
    слайда деки. Если python-pptx свойство однажды добавит — тест упадёт, и
    руками по XML можно будет перестать лазить.
    """
    from pptx.dml.fill import FillFormat

    assert not hasattr(FillFormat, "transparency")


def test_alpha_lands_on_the_colour_and_not_on_the_fill(canvas):
    """<a:alpha> пишется внутрь <a:srgbClr>, в тысячных долях процента."""
    _b, slide = canvas
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Emu(100), Emu(100))
    shp.fill.solid()
    shp.fill.fore_color.rgb = convert.RGBColor(0x31, 0xC5, 0x6E)

    assert set_alpha(shp.fill, 0.14) is True
    assert f'<a:srgbClr val="31C56E"><a:alpha val="14000"/>' in _xml(shp)
    # У самой заливки атрибутов прозрачности в DrawingML нет вовсе — если
    # альфа окажется там, PowerPoint её просто не увидит.
    assert shp.element.find(f".//{A}solidFill").attrib == {}


def test_alpha_is_replaced_and_not_stacked(canvas):
    """Второй вызов переписывает альфу, а не добавляет вторую.

    Два <a:alpha> в одном цвете — невалидный XML (в a:srgbClr группа
    EG_ColorTransform идёт как выбор), и PowerPoint предложит восстановить
    файл. Повторный вызов не выдумка: цвет и прозрачность у штрихов фона
    ставятся в цикле по одной и той же фигуре-заготовке.
    """
    _b, slide = canvas
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Emu(100), Emu(100))
    shp.fill.solid()
    shp.fill.fore_color.rgb = convert.RGBColor(0, 0, 0)

    set_alpha(shp.fill, 0.5)
    set_alpha(shp.fill, 0.25)
    assert len(shp.element.findall(f".//{A}srgbClr/{A}alpha")) == 1
    assert _alpha_of(shp) == 25000


def test_opaque_and_unpaintable_fills_write_nothing(canvas):
    """1.0 и заливка без цвета -> False и ни одного узла.

    Непрозрачность по умолчанию в DrawingML и так 100%, и лишний
    `<a:alpha val="100000"/>` только раздувал бы файл на каждой фигуре.
    """
    _b, slide = canvas
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Emu(100), Emu(100))
    shp.fill.solid()
    shp.fill.fore_color.rgb = convert.RGBColor(0, 0, 0)
    assert set_alpha(shp.fill, 1.0) is False
    assert _alpha_of(shp) is None

    bare = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Emu(100), Emu(100))
    bare.fill.background()                 # noFill: красить нечего
    assert set_alpha(bare.fill, 0.5) is False


def test_element_opacity_multiplies_the_colour_alpha(canvas):
    """opacity элемента и альфа цвета — два множителя, в файле одна альфа."""
    b, slide = canvas
    st = {"backgroundColor": "rgba(38, 208, 124, 0.4)", "opacity": "0.5"}
    shp = b._rect(slide, {"x": 0, "y": 0, "w": 100, "h": 50}, st)
    assert _alpha_of(shp) == 20000         # 0.4 * 0.5


# -------------------------------------------------------------------- рамки
def _borders(**px) -> dict:
    """Стиль с рамками: _borders(top=1, left=1) -> ширины и цвета как в палитре."""
    st = {f"border{s.capitalize()}Width": f"{px.get(s, 0)}px"
          for s in ("top", "bottom", "left", "right")}
    st["borderTopColor"] = st["borderLeftColor"] = "rgb(38, 208, 124)"
    return st


def test_only_the_sides_that_have_width_become_borders(canvas):
    """Ровно тот случай, из-за которого разделитель строк стал прямоугольником."""
    b, _slide = canvas
    assert set(b._sides(_borders(top=1))) == {"top"}
    assert set(b._sides(_borders(top=1, bottom=1, left=1, right=1))) == {
        "top", "bottom", "left", "right"}
    assert b._sides(_borders()) == {}
    # Цвета нет -> рамки нет: `border-width:1px` без цвета браузер не рисует.
    assert b._sides({"borderTopWidth": "1px"}) == {}


def test_one_border_becomes_a_line_on_that_side_only(canvas):
    """border-top -> отрезок по верхней кромке, а не обводка фигуры.

    Проверяем и ось отрезка: рамка в border-box лежит ВНУТРИ габарита, поэтому
    линия шириной 1px идёт по y+0.5, а не по y.
    """
    b, slide = canvas
    before = len(slide.shapes)
    it = {"x": 60, "y": 200, "w": 1800, "h": 120}

    assert b._rect(slide, it, _borders(top=1)) is None, "фигуры быть не должно"
    new = list(slide.shapes)[before:]
    assert len(new) == 1
    line = new[0]
    assert line.element.tag == P + "cxnSp", "сторона рамки — это линия (cxnSp)"
    assert line.left == Emu(round(60 * convert.EMU_PX))
    assert line.top == Emu(round(200.5 * convert.EMU_PX))
    assert line.width == Emu(round(1800 * convert.EMU_PX))
    assert line.height == 0
    assert line.line.width == Pt(0.5)      # 1 px слайда == 0.5 pt


def test_a_border_on_all_four_sides_stays_one_outline(canvas):
    """Рамка по всему контуру — обводка фигуры, а не четыре линии.

    Иначе каждая плашка деки стоила бы четырёх лишних объектов, а
    редактировать её пришлось бы по стороне.
    """
    b, slide = canvas
    before = len(slide.shapes)
    shp = b._rect(slide, {"x": 0, "y": 0, "w": 100, "h": 50},
                  _borders(top=1, bottom=1, left=1, right=1))

    assert len(list(slide.shapes)[before:]) == 1
    assert shp.element.tag == P + "sp"
    assert shp.element.find(f".//{A}ln/{A}solidFill") is not None
    assert shp.line.color.rgb == convert.RGBColor(38, 208, 124)


def test_two_sides_of_different_width_are_drawn_apart(canvas):
    """Неодинаковая рамка -> линии по сторонам, обводки у фигуры нет."""
    b, slide = canvas
    before = len(slide.shapes)
    st = _borders(top=1, bottom=4, left=1, right=1)
    st["backgroundColor"] = "rgb(0, 0, 0)"
    shp = b._rect(slide, {"x": 0, "y": 0, "w": 100, "h": 50}, st)

    # Сравнение по element, а не по объекту: slide.shapes каждый раз отдаёт
    # новую обёртку над тем же XML.
    lines = [s for s in list(slide.shapes)[before:]
             if s.element is not shp.element]
    assert len(lines) == 4, "одинаковой обводки нет — рисуем все четыре стороны"
    assert all(s.element.tag == P + "cxnSp" for s in lines)
    # У самой плашки обводки быть не должно: иначе рамка нарисуется дважды.
    assert shp.element.find(f".//{A}ln/{A}noFill") is not None


# --------------------------------------------------------------- фон полосами
def test_a_striped_background_becomes_real_bars(canvas):
    """«Шурешки» обложки: 90 штрихов по 2px с шагом 20 на полосе 1800px."""
    b, slide = canvas
    before = len(slide.shapes)
    it = {"x": 60, "y": 830, "w": 1800, "h": 183}

    assert b._stripes(slide, it, {"backgroundImage": HATCH_BG,
                                  "opacity": "1"}) == 90
    bars = list(slide.shapes)[before:]
    assert len(bars) == 90
    assert bars[0].left == Emu(round(60 * convert.EMU_PX))
    assert bars[1].left == Emu(round(80 * convert.EMU_PX))      # шаг 20
    assert bars[0].width == Emu(round(2 * convert.EMU_PX))      # штрих 2
    assert bars[0].height == Emu(round(183 * convert.EMU_PX))
    assert bars[-1].left == Emu(round((60 + 89 * 20) * convert.EMU_PX))
    assert bars[0].fill.fore_color.rgb == convert.RGBColor(38, 208, 124)
    assert _alpha_of(bars[0]) is None, "цвет непрозрачный — альфа лишняя"


def test_the_opacity_of_the_striped_element_dims_the_bars(canvas):
    """opacity:.5 у полосы -> альфа 50% у каждого штриха."""
    b, slide = canvas
    before = len(slide.shapes)
    b._stripes(slide, {"x": 0, "y": 0, "w": 100, "h": 10},
               {"backgroundImage": HATCH_BG, "opacity": "0.5"})
    bars = list(slide.shapes)[before:]
    assert bars and all(_alpha_of(s) == 50000 for s in bars)


def test_a_background_we_cannot_draw_is_counted_not_dropped_silently(canvas):
    """Незнакомый фон не переносим, но он попадает в отчёт.

    Молчаливая потеря однажды съела 40% слайдов деки, поэтому «не нарисовали»
    обязано быть видно: иначе новый градиент в шаблоне исчезнет со слайда и
    никто об этом не узнает до снимков из PowerPoint.
    """
    b, slide = canvas
    b.stats["dropped"].clear()
    assert b._background(slide, {"x": 0, "y": 0, "w": 100, "h": 100},
                         {"backgroundImage": "linear-gradient(red, blue)"}) == 0
    assert b.stats["dropped"] == {"linear-gradient": 1}

    # «Нет фона» — не потеря, в отчёт попадать не должно.
    b._background(slide, {"x": 0, "y": 0, "w": 100, "h": 100},
                  {"backgroundImage": "none"})
    assert b.stats["dropped"] == {"linear-gradient": 1}


def test_too_many_bars_are_refused_but_still_reported(canvas):
    """Защита по числу штрихов: 1px через 1px на всю ширину слайда — не рисуем.

    Предел стоит потому, что цена линейна: шаг в пиксель на полосе 1920 даёт
    почти две тысячи фигур на одном слайде, и это уже не редактируемый слайд, а
    неоткрываемый файл. Отказ попадает в отчёт — молчать нельзя и тут.
    """
    b, slide = canvas
    b.stats["dropped"].clear()
    dense = HATCH_BG.replace(" 20px)", " 1px)")
    before = len(slide.shapes)
    assert b._background(slide, {"x": 0, "y": 0, "w": 1920, "h": 10},
                         {"backgroundImage": dense}) == 0
    assert len(slide.shapes) == before
    assert b.stats["dropped"] == {"repeating-linear-gradient": 1}


# ------------------------------------------------------------- точечный фон
def _dot_style(step: int = 30, opacity: str = "1") -> dict:
    return {"backgroundImage": DOTS_BG, "backgroundSize": f"{step}px {step}px",
            "opacity": opacity}


def _squares(shape) -> list[tuple[int, int]]:
    """Центры квадратов контура: moveTo — левый верхний угол, второй lnTo —
    противоположный, их полусумма и есть узел сетки (целая: 2cx делится на 2).
    """
    out = []
    for path in shape.element.iter(f"{A}path"):
        for mv in path.findall(f"{A}moveTo"):
            p0 = mv.find(f"{A}pt")
            p2 = mv.getnext().getnext().find(f"{A}pt")
            out.append(((int(p0.get("x")) + int(p2.get("x"))) // 2,
                        (int(p0.get("y")) + int(p2.get("y"))) // 2))
    return out


def test_a_dotted_background_becomes_one_vector_shape(canvas):
    """Сетка точек — ОДНА фигура с квадратом в каждом узле, а не тысяча фигур.

    Отдельными фигурами это было бы 21612 штук на тёмную деку — файл, который
    PowerPoint открывает, но в котором уже нельзя работать. В CSS это тоже один
    элемент с фоновой картинкой, так что одна фигура — ещё и ближе к оригиналу.
    """
    b, slide = canvas
    before = len(slide.shapes)
    # 100x60 с шагом 30: узлы 20/50/80 по X и 20/50 по Y.
    assert b._dots(slide, {"x": 0, "y": 0, "w": 100, "h": 60}, _dot_style()) == 6
    assert len(slide.shapes) == before + 1

    shp = list(slide.shapes)[-1]
    assert "<a:custGeom" in _xml(shp), "форма должна быть собственной"
    assert "<a:prstGeom" not in _xml(shp), "готовая форма обязана уйти"
    assert len(_squares(shp)) == 6


def test_the_dots_are_squares_not_circles(canvas):
    """Узел сетки — квадрат со стороной в диаметр точки, а не круг.

    Фирменный паттерн Cloud.ru пиксельный. В CSS точка нарисована
    radial-gradient(circle...) — на трёх пикселях в браузере круг от квадрата
    не отличить, — но в PPTX точки превращаются в вектор, который разглядывают
    вблизи, и там узел обязан быть квадратным. Расхождение с разметкой
    сознательное и одностороннее (см. комментарий в _dots).
    """
    b, slide = canvas
    b._dots(slide, {"x": 0, "y": 0, "w": 100, "h": 60}, _dot_style())
    shp = list(slide.shapes)[-1]
    assert "<a:arcTo" not in _xml(shp), "узлы должны быть квадратами, не дугами"
    side = 2 * round(1.5 * convert.EMU_PX)        # диаметр точки из DOTS_BG
    for path in shp.element.iter(f"{A}path"):
        for mv in path.findall(f"{A}moveTo"):
            pts, node = [mv.find(f"{A}pt")], mv.getnext()
            while node is not None and node.tag == f"{A}lnTo":
                pts.append(node.find(f"{A}pt"))
                node = node.getnext()
            assert node is not None and node.tag == f"{A}close"
            xs = [int(p.get("x")) for p in pts]
            ys = [int(p.get("y")) for p in pts]
            assert len(pts) == 4, "у квадрата четыре угла"
            assert max(xs) - min(xs) == side and max(ys) - min(ys) == side


def test_the_dot_grid_sits_where_the_browser_puts_it(canvas):
    """Центр первой точки — сдвиг фона плюс половина плитки (замерено в Chrome).

    У .kpi-dots (рамка 60,600) снимок браузера даёт точку в (80,620), то есть
    +20 от края при шаге 30. Ошибись здесь на половину плитки — вся сетка
    поедет, а fidelity этого не увидит: рамка-то стоит на месте.
    """
    b, slide = canvas
    b._dots(slide, {"x": 0, "y": 0, "w": 100, "h": 60}, _dot_style(step=30))
    xs = sorted({x for x, _ in _squares(list(slide.shapes)[-1])})
    ys = sorted({y for _, y in _squares(list(slide.shapes)[-1])})
    assert [x / convert.EMU_PX for x in xs] == [20, 50, 80]
    assert [y / convert.EMU_PX for y in ys] == [20, 50]


def test_the_dot_offset_in_the_css_has_not_moved():
    """Страховка связи: сдвиг сетки лежит в коде, потому что его нет в замере.

    backgroundPosition в снимаемый набор свойств не входит, так что _DOT_ORIGIN
    — единственная копия значения из deck.css. Если правило в CSS поменяют,
    падать должен тест, а не точки на слайде.
    """
    css = Path("htmlslides/engine/deck.css").read_text(encoding="utf-8")
    rule = css.split(".pattern-dots{", 1)[1].split("}", 1)[0]
    assert f"background-position:{convert._DOT_ORIGIN:.0f}px " \
           f"{convert._DOT_ORIGIN:.0f}px" in rule


def test_the_opacity_of_the_dotted_element_dims_the_dots(canvas):
    """opacity:.5 у .kpi-dots — половинная альфа, как у штрихов и текста."""
    b, slide = canvas
    b._dots(slide, {"x": 0, "y": 0, "w": 100, "h": 60}, _dot_style(opacity="0.5"))
    assert _alpha_of(list(slide.shapes)[-1]) == 50000


def test_too_many_dots_are_refused_but_still_reported(canvas):
    """Защита по числу точек: шаг в пару пикселей — полмиллиона квадратиков.

    Предел нужен не для сегодняшней деки (её максимум — 4368 точек в области
    .contacts-dots светлой темы), а против шага, при котором один контур
    раздувает файл так, что PowerPoint его не открывает. Отказ попадает в
    отчёт. Предел обязан быть ВЫШЕ максимума деки: предыдущие 4000 были сняты
    с тёмной темы и молча резали контакты светлой.
    """
    b, slide = canvas
    b.stats["dropped"].clear()
    before = len(slide.shapes)
    assert b._background(slide, {"x": 0, "y": 0, "w": 1920, "h": 1080},
                         _dot_style(step=2)) == 0
    assert len(slide.shapes) == before
    assert b.stats["dropped"] == {"radial-gradient": 1}


# ------------------------------------------------------------- перенос строк
def _style(size: int = 30, weight: str = "400") -> dict:
    """Стиль нашим шрифтом — иначе измерителю нечем мерить."""
    return {"fontFamily": '"SB Sans Text", Verdana, sans-serif',
            "fontWeight": weight, "fontSize": f"{size}px",
            "letterSpacing": "normal", "lineHeight": f"{size * 1.15}px",
            "textTransform": "none"}


def test_a_frame_exactly_as_wide_as_its_text_is_not_wrapped():
    """Рамка в размер текста -> word_wrap выключен.

    Тот самый случай: таких рамок в деке 136 из 678, браузер подогнал их ширину
    под текст (inline-block, max-content), и запас там нулевой. Берём ширину
    ровно по измерителю — это и есть «ширина рамки равна ширине текста».
    """
    from htmlslides.topptx import textfit

    st = _style()
    exact = textfit.width("Аудит", textfit.Style(st))
    assert convert._wraps({"w": exact}, st, "Аудит") is False
    # И с запасом в пиксель — тоже одна строка, переносить нечего.
    assert convert._wraps({"w": exact + 1}, st, "Аудит") is False


def test_text_the_browser_did_wrap_is_left_wrappable():
    """Где перенёс браузер, переносит и PowerPoint — рамку рвать не надо."""
    st = _style()
    long = "Аудит инфраструктуры и сопровождение миграции в облако"
    assert convert._wraps({"w": 300}, st, long) is True


def test_padding_and_border_are_taken_off_the_measuring_strip():
    """Мерим content-box: рамка и padding тексту не принадлежат.

    Без этого полоса шире настоящей на 2px у .fg-cell, и слово, которое в
    браузере уехало на вторую строку, у нас осталось бы на первой.
    """
    from htmlslides.topptx import textfit

    st = _style()
    text = "Аудит инфраструктуры"
    exact = textfit.width(text, textfit.Style(st))
    padded = dict(st, paddingLeft="16px", paddingRight="16px",
                  borderLeftWidth="1px", borderRightWidth="1px")
    assert convert._wraps({"w": exact + 10}, st, text) is False
    assert convert._wraps({"w": exact + 10}, padded, text) is True


def test_the_reflowed_font_size_is_the_one_measured():
    """Кегль из строки (`fs`) важнее палитрового: рыба и слайд тут расходятся.

    У «цифр-героев» stats-row размер выставляет deck.js по длине значения
    (reflow.autofit_stats). Померив палитровым кеглем, мы бы считали строки для
    другого текста — и выключили перенос там, где он нужен, или наоборот.
    """
    from htmlslides.topptx import textfit

    st = _style(30)
    text = "Аудит инфраструктуры"
    narrow = textfit.width(text, textfit.Style(_style(30))) * 0.6
    assert convert._wraps({"w": narrow}, st, text) is True
    # Тот же текст тем же кеглем, но reflow уменьшил его вдвое — влезает.
    assert convert._wraps({"w": narrow, "fs": 15}, st, text) is False


def test_an_unfamiliar_font_is_left_to_powerpoint():
    """Чужой шрифт мерить нечем — не мешаем PowerPoint переносить сам."""
    st = dict(_style(), fontFamily="Arial, sans-serif")
    assert convert._wraps({"w": 10}, st, "Аудит") is True


# --------------------------------------------------------------- начертания
def test_weight_600_is_the_semibold_face_not_the_bold_flag():
    """font-weight:600 -> имя начертания «… Semibold», и НИКАКОГО флага b=1.

    Флаг жирности в PPTX один и означает вес 700: с ним PowerPoint на машине
    с полным семейством SB Sans берёт настоящий Bold, и вся дека выходит
    жирнее макета — ровно это и увидели на проверочной деке. Вес 600 несёт
    само имя: в name table шрифта начертание называется «Semibold» (так, слитно
    и с маленькой b — см. docstring face_for), и полное имя находят и macOS,
    и Windows.
    """
    from htmlslides.topptx.svgshapes import face_for

    assert face_for("SB Sans Display", "600") == (
        "SB Sans Display Semibold", False)
    assert face_for("SB Sans Text", "400") == ("SB Sans Text", False)
    # Настоящей жирности (700) в деке нет, но правило обязано её не съесть.
    assert face_for("SB Sans Text", "bold") == ("SB Sans Text", True)
    assert face_for("SB Sans Text", 700.0) == ("SB Sans Text", True)


def test_the_semibold_run_carries_the_face_name(canvas):
    """Ран текста веса 600 уходит в файл с именем начертания, без флага."""
    b, slide = canvas
    box = b._text(slide, {"x": 0, "y": 0, "w": 400, "h": 50},
                  _style(weight="600"), "Выручка")
    run = box.text_frame.paragraphs[0].runs[0]
    assert run.font.name == "SB Sans Text Semibold"
    assert not run.font.bold


# ------------------------------------------------------- и то же на деке
@pytest.fixture(scope="module")
def deck(tmp_path_factory):
    """Эталонная дека -> PPTX. Один раз на модуль: это полный проход Jinja."""
    from htmlslides.topptx import fidelity

    src = tmp_path_factory.mktemp("paint") / "deck.html"
    src.write_text(fidelity.build_deck("dark"), encoding="utf-8")
    out = src.with_suffix(".pptx")
    stats = convert.convert(src, out, theme="dark")
    return zipfile.ZipFile(out), stats


def test_a_deck_passed_as_markup_converts_like_one_passed_as_a_file(tmp_path):
    """Деку можно передать строкой, а не только путём, — и результат тот же.

    Это не удобство, а рабочий вход вебаппа: дека у него уже в памяти, писать
    её на диск незачем. Работало это только на коротких обрывках: решение
    «путь или разметка» принимал Path.exists(), а он на 350 KB не отвечает
    «нет», а падает с OSError «File name too long».
    """
    from htmlslides.topptx import fidelity

    html = fidelity.build_deck("dark")
    assert len(html) > 100_000, "дека вдруг стала короткой — проверка ослабла"
    a = convert.convert(html, tmp_path / "s.pptx", theme="dark")
    src = tmp_path / "deck.html"
    src.write_text(html, encoding="utf-8")
    b = convert.convert(src, tmp_path / "f.pptx", theme="dark")
    assert (a["slides"], a["shapes"], a["dot"]) == (b["slides"], b["shapes"],
                                                    b["dot"])


def test_a_path_that_does_not_exist_says_so(tmp_path):
    """Опечатка в пути — это «нет файла», а не «в деке нет слайдов».

    Короткая строка без тегов задумывалась путём, и отвечать на неё разбором
    как разметки значит отвечать не на тот вопрос.
    """
    with pytest.raises(convert.ConvertError, match="нет файла деки"):
        convert.convert(str(tmp_path / "нету.html"), tmp_path / "o.pptx")


def _slides(deck) -> list[etree._Element]:
    """Слайды в порядке деки — по номеру в имени части, а не по порядку в zip."""
    names = sorted((n for n in deck[0].namelist()
                    if n.startswith("ppt/slides/slide")),
                   key=lambda n: int("".join(c for c in n if c.isdigit())))
    return [etree.fromstring(deck[0].read(n)) for n in names]


def test_the_chart_wash_arrives_translucent(deck):
    """fill-opacity .14 из разметки -> <a:alpha val="14000"> в файле.

    Сквозная проверка проводки: set_alpha сам по себе уже проверен выше, здесь
    важно, что его кто-то зовёт на заливке SVG. Тот самый зелёный wash
    line-chart, который приезжал сплошной заливкой.
    """
    alphas = [int(a.get("val")) for s in _slides(deck)
              for a in s.findall(f".//{A}solidFill/{A}srgbClr/{A}alpha")]
    assert 14000 in alphas


def test_row_separators_are_lines_and_not_rectangles(deck):
    """В деке есть линии-стороны, и каждая — с цветом и толщиной.

    Линия без цвета была бы хуже прямоугольника: разделитель просто исчез бы.
    """
    # Линии в деке бывают двух родов: сторона рамки (их считает stats["edge"])
    # и <line> внутри svg (те идут в svg_shapes). Отличаем по имени — оно же и
    # видно в «Области выделения» PowerPoint.
    lines = [c for s in _slides(deck) for c in s.iter(P + "cxnSp")
             if "рамка" in (c.find(f".//{P}cNvPr").get("name") or "")]
    assert deck[1]["edge"] == len(lines) > 0
    for ln in lines:
        assert ln.find(f".//{A}ln/{A}solidFill") is not None
        assert ln.find(f".//{A}ln").get("w"), "толщина линии не задана"


def test_the_cover_keeps_its_hatch(deck):
    """Обложка: 90 штрихов, и у каждого своё имя — объект, а не картинка."""
    cover = _slides(deck)[0]
    names = [c.get("name") for c in cover.iter(P + "cNvPr")]
    bars = [n for n in names if n and "штрих" in n]
    assert len(bars) == 90
    assert len(set(bars)) == 90, "имена штрихов должны быть различимы"


def test_single_line_text_is_not_rewrapped_by_powerpoint(deck):
    """«СЕТЬ» и «Аудит» приезжают с wrap="none", перенесённый текст — со square.

    Проверка на обе стороны нужна: выключить перенос везде так же неверно, как
    включить. Текст, который в браузере занял две строки, в PPTX обязан
    переноситься — иначе он уедет одной строкой за край слайда.
    """
    tight, loose = [], []
    for s in _slides(deck):
        for tx in s.iter(P + "txBody"):
            body = tx.find(A + "bodyPr")
            words = "".join(t.text or "" for t in tx.iter(A + "t"))
            (tight if body.get("wrap") == "none" else loose).append(words)

    assert tight, "перенос не выключен ни у одной рамки"
    assert loose, "перенос выключен у всех — многострочный текст уедет за край"
    assert any(" " not in w.strip() and w.strip() for w in tight), (
        "одиночное слово в рамке по себе обязано остаться одной строкой")
    # Самая длинная надпись деки — заведомо не одна строка, переносить её надо.
    assert max(tight + loose, key=len) in loose
